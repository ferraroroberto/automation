# requirements: public
import logging
import pandas as pd
import warnings
from openpyxl import load_workbook

# requirements: custom functions
from utils import read_params_from_txt_file, DEFAULT_PARAMS_FILE

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
log = logging.getLogger(__name__)

# Suppress openpyxl warnings
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl.worksheet._reader")

def read_database_editorial(excel_path):
    df = pd.read_excel(excel_path)
    return df[df["ind_editorial"] == 1]

def editorial_databases(databases_to_editorial):
    log.info("Databases to process: %d", len(databases_to_editorial))

    for _, database in databases_to_editorial.iterrows():
        log.info("Processing excel copy to editorial calendar for database: %s", database['name'])
        editorial_name = database["editorial_name"]
        output_path_clean = database["output_path_clean"]

        try:
            # Load source workbook and destination workbook.
            src_wb = load_workbook(filename=output_path_clean)
            dest_wb = load_workbook(filename=editorial_excel_path)

            # Get source sheet
            src_sheet = src_wb.active

            # Get or create sheet in the destination workbook
            if editorial_name in dest_wb.sheetnames:
                log.info('Sheet "%s" already exists in the destination file.', editorial_name)
                dest_sheet = dest_wb[editorial_name]
                for row in dest_sheet.iter_rows():
                    for cell in row:
                        cell.value = None
            else:
                log.info('Sheet "%s" does not yet exist in the destination file.', editorial_name)
                dest_sheet = dest_wb.create_sheet(title=editorial_name)

            # Initialize a counter for the rows
            rows_copied = 0

            # Copy the data from source sheet to destination sheet
            for row in src_sheet.iter_rows():
                for cell in row:
                    dest_sheet[cell.coordinate].value = cell.value
                rows_copied += 1

            # Save the changes in the destination workbook
            dest_wb.save(editorial_excel_path)

            log.info('Sheet "%s" copied to the destination file. Copied %d rows.', editorial_name, rows_copied)

        except Exception as e:
            log.error("Error copying the sheet: %s", e)

        log.info("Processed database '%s' saved to %s as the sheet %s", database['name'], editorial_excel_path, editorial_name)

# Main execution

params_file_path = DEFAULT_PARAMS_FILE
params = read_params_from_txt_file(params_file_path)

excel_path = params['excel_path']
editorial_excel_path = params['editorial_excel_path']

# this part is from here "notion-09-pass-to-editorial" > https://chat.openai.com/c/1279d9a7-829a-4e0d-a608-e75252e501d3
databases_to_editorial = read_database_editorial(excel_path)
editorial_databases(databases_to_editorial)