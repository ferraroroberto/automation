#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Storage Management Module

This module handles all file operations including Excel file loading, saving, and data management
for the LinkedIn image search application.

Author: Roberto
Date: March 2025
"""

import logging
import numpy as np
import pandas as pd
from pathlib import Path
import time
from typing import Dict, Tuple

try:
    import openpyxl
    import openpyxl.utils
except ImportError:
    openpyxl = None

logger = logging.getLogger(__name__)


class StorageManager:
    """Manages storage operations for metadata, results, and API history."""
    
    def __init__(self, metadata_folder: Path):
        """
        Initialize storage manager.
        
        Args:
            metadata_folder (Path): Folder for storing Excel files
        """
        self.metadata_folder = metadata_folder
        self.metadata_folder.mkdir(parents=True, exist_ok=True)
        
    def setup_folders_and_files(self, images_folder: Path, metadata_file: str, 
                               results_file: str, api_history_file: str) -> Tuple[Path, Path, Path, Path]:
        """
        Setup folders and file paths.
        
        Args:
            images_folder (Path): Path to images folder
            metadata_file (str): Name of metadata Excel file
            results_file (str): Name of results Excel file
            api_history_file (str): Name of API history Excel file
            
        Returns:
            Tuple[Path, Path, Path, Path]: Paths for images folder, metadata file, results file, and API history file
            
        Raises:
            FileNotFoundError: If images folder doesn't exist
        """
        if not images_folder.exists():
            logger.error(f"❌ Error: Images folder '{images_folder}' does not exist.")
            raise FileNotFoundError(f"Images folder not found: {images_folder}")
        
        logger.info(f"✅ Images folder exists: {images_folder}")
        logger.info(f"✅ Metadata folder exists/created: {self.metadata_folder}")
        
        # Define paths for metadata, results, and API history files
        metadata_path = self.metadata_folder / metadata_file
        results_path = self.metadata_folder / results_file
        api_history_path = self.metadata_folder / api_history_file
        
        return images_folder, metadata_path, results_path, api_history_path
    
    def load_metadata(self, metadata_path: Path) -> pd.DataFrame:
        """
        Load metadata file if it exists, or create an empty DataFrame.
        
        Args:
            metadata_path (Path): Path to the metadata file
            
        Returns:
            pd.DataFrame: Metadata DataFrame with images and their Imgur URLs
        """
        # Define the essential columns we require
        essential_columns = ['filename', 'last_processed_date', 'imgur_url']
        
        # Define count columns we'll add/update
        count_columns = [
            'total_links', 'exact_match_count', 'similar_match_count',
            'linkedin_count', 'instagram_count', 'twitter_count', 
            'facebook_count', 'pinterest_count'
        ]
        
        if metadata_path.exists():
            try:
                logger.info(f"📊 Loading existing metadata from {metadata_path}")
                metadata_df = pd.read_excel(metadata_path)
                
                # Validate essential columns
                if not all(col in metadata_df.columns for col in essential_columns):
                    logger.warning("⚠️ Metadata file is missing essential columns. Creating backup and initializing new file.")
                    self._create_backup(metadata_path, metadata_df)
                    metadata_df = pd.DataFrame(columns=essential_columns)
                else:
                    # Add any missing count columns (for backward compatibility)
                    for col in count_columns:
                        if col not in metadata_df.columns:
                            insert_pos = len(essential_columns)
                            metadata_df.insert(insert_pos, col, 0)
                            logger.info(f"📊 Added missing column '{col}' to metadata")
                
                return metadata_df
            except Exception as e:
                logger.error(f"❌ Error loading metadata file: {e}")
                logger.info("🔄 Initializing new metadata file")
                return self._create_empty_metadata_df(essential_columns, count_columns)
        else:
            logger.info("📊 No existing metadata found. Initializing new metadata file.")
            return self._create_empty_metadata_df(essential_columns, count_columns)
    
    def load_results_database(self, results_path: Path) -> pd.DataFrame:
        """
        Load existing results database if it exists, or create an empty DataFrame.
        
        Args:
            results_path (Path): Path to the results database file
            
        Returns:
            pd.DataFrame: Results database with all previous search results
        """
        # Define the required columns for the results database
        required_columns = ['local_image', 'found_link']
        
        # Define standard columns that should be present
        standard_columns = [
            'local_image', 'uploaded_url', 'found_link', 
            'title', 'duplicate', 'match_type', 'source', 
            'post_date', 'search_date', 'order'
        ]
        
        if results_path.exists():
            try:
                logger.info(f"📊 Loading existing results database from {results_path}")
                results_df = pd.read_excel(results_path)
                
                # Validate that we have at least the required columns
                if not all(col in results_df.columns for col in required_columns):
                    logger.warning("⚠️ Results file is missing required columns. Creating backup and initializing new file.")
                    self._create_backup(results_path, results_df)
                    results_df = pd.DataFrame(columns=standard_columns)
                
                # Handle the transition from 'snippet' to 'duplicate' column
                results_df = self._handle_snippet_to_duplicate_migration(results_df)
                
                return results_df
            except Exception as e:
                logger.error(f"❌ Error loading results file: {e}")
                logger.info("🔄 Initializing new results database")
                return pd.DataFrame(columns=standard_columns)
        else:
            logger.info("📊 No existing results database found. Initializing new database.")
            return pd.DataFrame(columns=standard_columns)
    
    def load_api_history_database(self, history_path: Path) -> pd.DataFrame:
        """
        Load existing API history database if it exists, or create an empty DataFrame.
        
        Args:
            history_path (Path): Path to the API history database file
            
        Returns:
            pd.DataFrame: API history database with all previous API calls
        """
        # Define the required columns for API history
        required_columns = ['search_type', 'local_image', 'imgur_url', 'search_date']
        
        # Define standard columns for API history
        standard_columns = [
            'search_type', 'local_image', 'imgur_url', 'search_date',
            'api_id', 'api_status', 'json_endpoint', 'created_at', 'processed_at',
            'total_time_taken', 'engine', 'url', 'page_token', 'raw_json',
            'search_engine_query', 'raw_html_file'
        ]
        
        # Optional columns that might be present depending on API response
        optional_columns = [
            'user', 'requester_ip', 'rtt', 'device', 'source_user_agent', 
            'used_user_agent', 'proxy_provider', 'playground_link'
        ]
        
        if history_path.exists():
            try:
                logger.info(f"📊 Loading existing API history database from {history_path}")
                history_df = pd.read_excel(history_path)
                
                # Validate that we have at least required columns
                if not all(col in history_df.columns for col in required_columns):
                    logger.warning("⚠️ API history file is missing required columns. Creating backup and initializing new file.")
                    self._create_backup(history_path, history_df)
                    all_columns = standard_columns + optional_columns
                    history_df = pd.DataFrame(columns=all_columns)
                
                return history_df
            except Exception as e:
                logger.error(f"❌ Error loading API history file: {e}")
                logger.info("🔄 Initializing new API history database")
                all_columns = standard_columns + optional_columns
                return pd.DataFrame(columns=all_columns)
        else:
            logger.info("📊 No existing API history database found. Initializing new database.")
            all_columns = standard_columns + optional_columns
            return pd.DataFrame(columns=all_columns)
    
    def save_excel_with_retry(self, df: pd.DataFrame, file_path: Path, 
                             max_retries: int = 5, retry_delay: int = 5) -> bool:
        """
        Save an Excel file with retry logic and column width preservation.
        
        Args:
            df (pd.DataFrame): DataFrame to save
            file_path (Path): Path to save the file
            max_retries (int): Maximum number of retry attempts
            retry_delay (int): Delay in seconds between retries
            
        Returns:
            bool: True if successful, False otherwise
        """
        # Read column widths from existing file if it exists
        column_widths = self._get_excel_column_widths(file_path)
        
        # Preserve custom columns from existing file
        df = self._preserve_custom_columns(df, file_path)
        
        # Try to save the file with retries
        for attempt in range(max_retries):
            try:
                # Save the DataFrame to Excel
                with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                    df.to_excel(writer, index=False)
                    
                    # Apply column widths if we have them and openpyxl is available
                    if column_widths and openpyxl:
                        worksheet = writer.sheets['Sheet1']
                        for col_idx, width in column_widths.items():
                            if col_idx < len(df.columns):
                                col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
                                worksheet.column_dimensions[col_letter].width = width
                
                logger.info(f"✅ Saved file with preserved formatting: {file_path}")
                return True
            except PermissionError:
                if attempt < max_retries - 1:
                    logger.warning(f"❌ Excel file '{file_path}' is open. Please close it. Retrying in {retry_delay} seconds...")
                    time.sleep(retry_delay)
                else:
                    logger.error(f"❌ Failed to save after {max_retries} attempts: {file_path}")
                    return False
            except Exception as e:
                logger.error(f"❌ Error saving file {file_path}: {str(e)}")
                return False
        
        return False
    
    def _create_empty_metadata_df(self, essential_columns: list, count_columns: list) -> pd.DataFrame:
        """Create empty metadata DataFrame with all required columns."""
        df = pd.DataFrame(columns=essential_columns)
        for col in count_columns:
            df[col] = 0
        return df
    
    def _create_backup(self, file_path: Path, df: pd.DataFrame) -> None:
        """Create a backup of an existing file."""
        backup_path = file_path.with_suffix(f".bak.{int(time.time())}.xlsx")
        df.to_excel(backup_path, index=False)
        logger.info(f"📋 Created backup: {backup_path}")
    
    def _handle_snippet_to_duplicate_migration(self, results_df: pd.DataFrame) -> pd.DataFrame:
        """Handle migration from 'snippet' to 'duplicate' column."""
        if 'snippet' in results_df.columns and 'duplicate' not in results_df.columns:
            logger.info("🔄 Converting 'snippet' column to 'duplicate' column")
            results_df = results_df.rename(columns={'snippet': 'duplicate'})
            results_df['duplicate'] = 0
        elif 'snippet' in results_df.columns and 'duplicate' in results_df.columns:
            logger.info("🔄 Dropping 'snippet' column as 'duplicate' already exists")
            results_df = results_df.drop(columns=['snippet'])
        elif 'duplicate' not in results_df.columns:
            logger.info("🔄 Adding 'duplicate' column")
            results_df['duplicate'] = 0
        
        return results_df
    
    def _get_excel_column_widths(self, file_path: Path) -> Dict[int, float]:
        """Read column widths from an existing Excel file."""
        if not file_path.exists() or not openpyxl:
            return {}
            
        try:
            wb = openpyxl.load_workbook(file_path)
            sheet = wb.active
            
            col_widths = {}
            for i, column in enumerate(sheet.columns):
                col_letter = openpyxl.utils.get_column_letter(i + 1)
                if hasattr(sheet.column_dimensions[col_letter], 'width'):
                    col_widths[i] = sheet.column_dimensions[col_letter].width
            
            return col_widths
        except Exception as e:
            logger.warning(f"⚠️ Could not read column widths from {file_path}: {e}")
            return {}
    
    def _preserve_custom_columns(self, df: pd.DataFrame, file_path: Path) -> pd.DataFrame:
        """Preserve custom columns from existing file."""
        if not file_path.exists():
            return df
            
        try:
            existing_df = pd.read_excel(file_path)
        except Exception as e:
            logger.warning(f"⚠️ Could not read existing file for custom columns: {e}")
            return df
        
        # Find columns in the existing file that aren't in our DataFrame
        custom_cols = [col for col in existing_df.columns if col not in df.columns]
        
        if not custom_cols:
            return df
            
        logger.info(f"📊 Preserving {len(custom_cols)} custom columns: {', '.join(custom_cols)}")
        
        # Determine key columns for mapping rows
        key_cols = self._get_key_columns(df)
        
        if not key_cols or not all(key in existing_df.columns for key in key_cols):
            logger.warning(f"⚠️ Could not determine key columns for {file_path}, skipping custom columns")
            return df
        
        # Add custom columns to our DataFrame
        for col in custom_cols:
            df[col] = np.nan
            
            # Copy values from existing DataFrame where keys match
            for _, row in existing_df.iterrows():
                filter_expr = True
                for key in key_cols:
                    filter_expr = filter_expr & (df[key] == row[key])
                
                if filter_expr.any():
                    df.loc[filter_expr, col] = row[col]
        
        return df
    
    def _get_key_columns(self, df: pd.DataFrame) -> list:
        """Determine key columns for row mapping based on DataFrame structure."""
        if 'filename' in df.columns:
            return ['filename']  # Metadata file
        elif 'local_image' in df.columns and 'found_link' in df.columns:
            return ['local_image', 'found_link']  # Results file
        elif 'search_type' in df.columns and 'local_image' in df.columns and 'search_date' in df.columns:
            return ['search_type', 'local_image', 'search_date']  # API history file
        else:
            return [] 